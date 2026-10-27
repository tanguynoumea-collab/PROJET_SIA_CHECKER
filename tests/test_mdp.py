"""Tableau MDP (mdp.py) et remplissage des tableaux Excel sans perte (xlsx_tables.py), sur le gabarit livré."""

import re
import warnings
import zipfile

import openpyxl
import pytest

from sia_check import mdp
from sia_check.xlsx_tables import Formula, TemplateError, Workbook

TEMPLATE = mdp.DEFAULT_TEMPLATE


def _load(path, **kw):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")   # openpyxl : « Data Validation extension is not supported »
        return openpyxl.load_workbook(path, **kw)


def _ext(path) -> dict[str, list[bytes]]:
    """Blocs extLst (validations x14) de chaque feuille."""
    with zipfile.ZipFile(path) as z:
        return {n: re.findall(rb"<extLst>.*</extLst>", z.read(n), re.S)
                for n in z.namelist() if n.startswith("xl/worksheets/sheet")}


@pytest.fixture
def filled(result, cfg, tmp_path):
    return mdp.write(result, cfg, tmp_path / "MDP-42-Synth.xlsx", project_id=42, project_name="Synth")


def test_template_is_blank():
    wb = Workbook(TEMPLATE)
    for name in ("SurfacesSIA4168", "SurfacesSIA41610", "SurfacesSIA41611", "Logements", "Tableau3"):
        assert all(v is None for r in wb.table_rows(name) for v in r.values()), name
    assert not any(r["Valeur"] for r in wb.table_rows("InfosProjet"))


def test_mdp_keeps_validations_pivot_and_lists(filled):
    assert _ext(filled) == _ext(TEMPLATE)                      # listes déroulantes intactes
    with zipfile.ZipFile(filled) as z:
        names = set(z.namelist())
    assert "xl/pivotTables/pivotTable1.xml" in names and "xl/calcChain.xml" not in names
    with zipfile.ZipFile(TEMPLATE) as z:
        tpl = {n: z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(filled) as z:
        assert z.read("xl/worksheets/sheet9.xml") == tpl["xl/worksheets/sheet9.xml"]   # Listes formulaire


def test_mdp_tables_content(filled, result, cfg):
    wb = _load(filled)
    brutes = wb["SurfacesBrutesVolumeSIA416"]
    nettes = wb["SurfacesNettesSIA416"]
    n_brutes = int(((result.surfaces.kind == "node") & result.surfaces.code.isin(["SP", "SEP"])).sum())
    assert brutes.tables["SurfacesSIA4168"].ref == f"A1:L{1 + n_brutes}"
    row = [c.value for c in brutes[2]]
    assert row[0] == "=InfosProjet!$C$2" and row[1] == "SP" and row[4] == "Hors-sol"
    m3 = [(r[9], r[11]) for r in brutes.iter_rows(min_row=2, values_only=True)]
    assert all(f.startswith("=SurfacesSIA4168[[#This Row],[m2]]") if h else f is None for h, f in m3)   # sans hauteur : vide
    cats = {r[1] for r in nettes.iter_rows(min_row=2, values_only=True)}
    assert {"SUP", "SED", "SUT"} <= cats and "SP" not in cats
    uspi = [r for r in wb["SurfacesUSPI"].iter_rows(min_row=2, values_only=True) if r[2] == "OK_N01_Apart-1"]
    assert {r[1]: r[3] for r in uspi}["Surface externe"] == pytest.approx(90 - 85)    # SE = SVN − SH
    info = {r[1]: r[2] for r in _load(filled, data_only=True)["InfosProjet"].iter_rows(min_row=2, values_only=True)}
    assert info["ID projet"] == 42 and info["Nom Projet HRS"] == "Synth" and info["Nbr total de logements"] >= 1


def test_cached_values_for_formulas(filled):
    ws = _load(filled, data_only=True)["SurfacesBrutesVolumeSIA416"]
    r = [c.value for c in ws[2]]
    assert r[0] == 42 and r[11] == pytest.approx(r[10] * r[9]) if r[9] else r[11] is None


def test_fill_table_moves_totals_row_and_keeps_styles(tmp_path):
    wb = Workbook(TEMPLATE)
    recs = [{"Partie de l'enveloppe": "Facade global", "Etage": f"N0{i}", "Surface Brutte": 10.0 * i} for i in range(5)]
    wb.fill_table("SurfacesSIA416", recs)
    out = wb.save(tmp_path / "t.xlsx")
    ws = _load(out)["SurfacesEnveloppe"]
    assert ws.tables["SurfacesSIA416"].ref == "A1:I7"            # 5 lignes + totaux
    assert ws["D6"].value == "N04" and ws["G6"].value == 40
    assert ws["G2"].number_format == ws["G6"].number_format


def test_filter_lifted_and_ref_formula_removed():
    with zipfile.ZipFile(TEMPLATE) as z:
        assert b"#REF!" not in z.read("xl/tables/table4.xml")
        assert b"filterColumn" not in z.read("xl/tables/table3.xml")


def test_unknown_column_and_table_raise():
    wb = Workbook(TEMPLATE)
    with pytest.raises(TemplateError):
        wb.fill_table("SurfacesSIA41611", [{"Colonne inconnue": 1}])
    with pytest.raises(TemplateError):
        wb.fill_table("Absent", [])


def test_formula_cached_string_and_inline_text(tmp_path):
    wb = Workbook(TEMPLATE)
    wb.fill_table("SurfacesSIA41611", [{"ID Projet": Formula("InfosProjet!$C$2", "P-7"), "Categorie USPI": "a < b & c",
                                        "Appartement_Unite": "X", "m2": 1.5}])
    out = wb.save(tmp_path / "t.xlsx")
    ws = _load(out, data_only=True)["SurfacesUSPI"]
    assert [c.value for c in ws[2]] == ["P-7", "a < b & c", "X", 1.5]


def test_compact_strings_drops_unused(tmp_path):
    wb = Workbook(TEMPLATE)
    out = wb.save(tmp_path / "t.xlsx", compact_strings=True)
    with zipfile.ZipFile(out) as z:
        n = z.read("xl/sharedStrings.xml").count(b"<si>")
    with zipfile.ZipFile(TEMPLATE) as z:
        assert n == z.read("xl/sharedStrings.xml").count(b"<si>")   # gabarit déjà compact : rien de perdu
    assert _load(out)["InfosProjet"]["B2"].value == "ID projet"


IDENTITE = {"ID projet": "1778", "Nom Projet HRS": "L'Orée", "Canton": "VD", "Commune": "Lausanne",
            "Categorie d'Ouvrage": "Santé", "Type d'intervention": "Nouvelle construction", "Phase du projet": "53"}


def _infos(path) -> dict:
    ws = _load(path, data_only=True)["InfosProjet"]
    return {r[1]: r[2] for r in ws.iter_rows(values_only=True) if r[1]}


def test_identite_in_exported_mdp(result, cfg, tmp_path):
    path = mdp.write(result, cfg, tmp_path / "MDP.xlsx", project_id=42, project_name="Synth", identite=IDENTITE)
    infos = _infos(path)
    assert infos["ID projet"] == 1778 and infos["Phase du projet"] == 53     # nombres, pour la jointure Power BI
    assert infos["Canton"] == "VD" and infos["Categorie d'Ouvrage"] == "Santé" and infos["Nom Projet HRS"] == "L'Orée"
    rows = list(_load(path, data_only=True)["SurfacesNettesSIA416"].iter_rows(min_row=2, max_row=2, values_only=True))
    assert rows[0][0] == 1778                                                # l'ID de la fiche relie les tables


def test_identite_written_into_existing_mdp_keeps_the_rest(filled):
    before = _ext(filled)
    mdp.ecrire_identite(filled, {"Canton": "GE", "Commune": "", "Categorie d'Ouvrage": "EMS"})
    infos = _infos(filled)
    assert infos["Canton"] == "GE" and infos["Categorie d'Ouvrage"] == "EMS"
    assert infos["Nom Projet HRS"] == "Synth" and infos["ID projet"] == 42   # valeur vide : rien d'effacé
    assert _ext(filled) == before                                            # listes déroulantes intactes


def test_listes_of_the_template():
    listes = mdp.listes()
    assert "Santé" in listes["Categorie d'Ouvrage"] and "x" not in listes["Categorie d'Ouvrage"]
    assert "VD" in listes["Canton"] and "Abreviation" not in listes["Canton"]
    assert listes["Phase du projet"][0] == "21"
