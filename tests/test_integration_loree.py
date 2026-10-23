"""Intégration sur le vrai projet L'Orée (ignoré si le projet Treegram n'est pas sur le poste).

Chemin surchargeable par la variable d'environnement SIA_LOREE_DIR.
"""

import json
import os
from pathlib import Path

import pytest

from sia_check import cli, reader, rules, treegram

LOREE = Path(os.environ.get(
    "SIA_LOREE_DIR",
    Path.home() / "OneDrive - HRS Real Estate AG/Documents/Treegram/Projects/L OREE",
))
TGM = LOREE / "L OREE.tgm"
IFC = LOREE / "Exports/L OREE - Export Surfaces.ifc"
BUILDINGS = {"A01", "B03", "B04", "E09", "G12", "H16", "I14", "I17", "I18"}

needs_tgm = pytest.mark.skipif(not TGM.is_file(), reason="projet L'Orée absent")


@pytest.fixture(scope="module")
def loree(request):
    from sia_check import config
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    df = reader.read(TGM, cfg)
    return df, rules.run(df, cfg)


def failing(res, rule):
    f = res.failing
    return {(r.building, r.level) for r in f[f.rule == rule].itertuples()}


@needs_tgm
def test_counts_match_viewer(loree):
    # Compteurs du viewer Treegram : 80 SP, 80 SEP (un total par niveau).
    df, _ = loree
    assert (df.category != "SFN").sum() == 6090
    sfn = df[df.category == "SFN"]  # façade de chaque niveau, même hauteur brute que la SP du niveau
    assert len(sfn) == 80 and len(sfn[["building", "level"]].drop_duplicates()) == 80
    assert sfn.gross_area.sum() == pytest.approx(30138.23, abs=0.01)
    nodes = df[df.kind == "node"].category.value_counts()
    assert nodes["SP"] == 80 and nodes["SEP"] == 80
    assert set(df.building) == BUILDINGS
    assert df[(df.kind == "node") & (df.category == "SP")].area.sum() == pytest.approx(56210.12, abs=0.01)


@needs_tgm
def test_known_anomaly_level_without_su(loree):
    # Anomalie connue « niveau avec SP sans SU » (I18 N00) : en norme stricte elle sort en R5 ;
    # avec la règle personnalisée P1 (SU = SUP + SUS + SUT), ses 284 m² de SUT forment sa surface utile.
    import yaml
    from sia_check import config
    df, res = loree
    assert ("I18", "N00") not in failing(res, "R5")
    raw = yaml.safe_load((Path(__file__).resolve().parents[1] / "config.yaml").read_text(encoding="utf-8"))
    raw.pop("custom_rules")
    assert ("I18", "N00") in failing(rules.run(df, config.from_dict(raw)), "R5")


@needs_tgm
def test_known_anomaly_ground_sep(loree):
    _, res = loree
    ground = {(b, "N00") for b in BUILDINGS}
    assert failing(res, "A1") == ground
    assert ground <= failing(res, "R3")


@needs_tgm
def test_known_anomaly_roof_sep(loree):
    _, res = loree
    roofs = {(b, "Toiture") for b in BUILDINGS}
    assert failing(res, "A2") == roofs
    assert roofs <= failing(res, "R3")


@needs_tgm
def test_uspi_chain_consistent(loree):
    _, res = loree
    u = res.residuals[res.residuals.rule.str.startswith("U")]
    assert len(u) > 1000 and (u.verdict == "OK").all()


@needs_tgm
def test_cli_exit_code(capsys):
    assert cli.main([str(TGM), "--config", str(Path(__file__).resolve().parents[1] / "config.yaml")]) == 1
    assert "A1" in capsys.readouterr().out


@pytest.mark.skipif(not IFC.is_file(), reason="export IFC L'Orée absent")
def test_promoter_ifc_only_unknown_codes():
    # L'export IFC Treegram porte la nomenclature promoteur : tout est inconnu, rien ne plante.
    # Exception : « ST » (surface taxable française) a le même code que ST SIA 416 §1 (surface de terrain).
    from sia_check import config
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    res = rules.run(reader.read(IFC, cfg), cfg)
    assert set(res.residuals.rule) == {"A3", "R0"}
    assert set(res.surfaces.category) == {"ST"}
    assert len(res.unknown) == 880 - 80


@needs_tgm
def test_treegram_exports_match_treegram(loree):
    # Comparé le 22.10.2026 aux exports Treegram 1778-SurfacesNettes / -SurfacesBrutes : mêmes lignes, mêmes
    # valeurs, même ordre ; s'y ajoutent les 26 SED et 47 SUT que l'export Treegram omet.
    from sia_check import config
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    _, res = loree
    n, b = treegram.nettes(res, cfg), treegram.brutes(res, cfg)
    assert len(n) == 3221 + 26 + 47 and len(b) == 160
    assert n["Catégorie SIA"].value_counts()[["SED", "SUT"]].tolist() == [26, 47]
    assert n["ID TGM"].iloc[0] == "H170::H2527" and n["ID TGM"].is_unique and n["ID TGM"].notna().all()
    assert b.Area.sum() == pytest.approx(74518.54, abs=0.01)
    assert b.Perimeter.sum() == pytest.approx(17233.96, abs=0.01)


@needs_tgm
def test_mdp_from_blank_template(loree, tmp_path):
    import warnings

    import openpyxl

    from sia_check import config, mdp
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    _, res = loree
    out = mdp.write(res, cfg, tmp_path / "MDP-1778-Loree.xlsx", project_id=1778, project_name="L'Orée")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(out)
    refs = {t: ws.tables[t].ref for ws in wb.worksheets for t in ws.tables}
    assert refs["SurfacesSIA4168"] == "A1:L152" and refs["SurfacesSIA41610"] == "A1:I3295"   # sans les 9 SP Toiture à 0 m²
    assert refs["Logements"] == "A1:K316"                       # 315 lots
    env = list(wb["SurfacesEnveloppe"].iter_rows(min_row=2, max_col=9, values_only=True))
    assert sum(1 for r in env if r[1] == mdp.FACADE) == 80
    info = {r[1]: r[2] for r in wb["InfosProjet"].iter_rows(min_row=2, values_only=True)}
    assert info["Nbr de bâtiments (lots)"] == 9 and info["Nbr total de logements"] == 315


@needs_tgm
def test_corrections_s1_s2(loree):
    # S2 : SEP des niveaux Toiture exclues ; S1 : 1 800 m² de la SEP B03 N00 (dalle sur sous-sol) découpés en SA.
    from sia_check import config, corrections
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    df, _ = loree
    journal = {"version": 1, "corrections": [
        {"id": "s2", "op": "exclure", "statut": "acceptée", "sujet": "S2",
         "cible": {"selection": {"category": "SEP", "level": "Toiture", "kind": "node"}}},
        {"id": "s1-b03", "op": "découper", "statut": "acceptée", "sujet": "S1", "vers": "SA", "aire": 1800.0,
         "cible": {"ids": ["B03_N00_SEP#total"]}},
    ]}
    applied, before, after = corrections.run(df, journal, cfg)
    assert applied.log.set_index("id").loc["s2", "surfaces"] == 9
    toit = lambda r: {x.rule for x in r.failing.itertuples() if x.level == "Toiture"}  # noqa: E731
    assert {"R3", "A2"} <= toit(before) and not toit(after) & {"R3", "A2"}
    r3 = lambda r: r.residuals.query("rule == 'R3' and building == 'B03' and level == 'N00'").diff_m2.iloc[0]  # noqa: E731
    assert r3(before) - r3(after) == pytest.approx(1800.0)
    assert not (after.residuals.query("rule == 'R4' and category == 'ST'")).size


@needs_tgm
def test_subjects_without_ai(loree):
    # Revue IA du 22.10.2026 : S1 (abords, B03 N00 ≈ 1 812 m²), S2 (toitures ≈ 7 279 m²), S6 (SC implicite).
    pytest.importorskip("shapely")
    from sia_check import config, corrections, geom, sujets, viewer
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    df, res = loree
    doc = viewer.build(res, cfg, geom.read_cache(LOREE), "L OREE", [str(TGM)])
    sj = {s["id"]: s for s in sujets.build(doc, cfg)["subjects"]}
    assert sj["toitures"]["m2"] == pytest.approx(7279, abs=5)
    b03 = next(c for c in sj["abords"]["corrections"] if c["id"] == "abords-B03-N00")
    assert b03["aire"] == pytest.approx(1812.6, abs=1) and b03["polygone"]
    assert sj["sc-implicite"]["action"] == "aucune" and len(sj["sc-implicite"]["points"]) > 50
    journal = sujets.journal({"subjects": list(sj.values())})
    for c in journal["corrections"]:
        c["statut"] = "acceptée"
    applied, before, after = corrections.run(df, journal, cfg)
    assert set(applied.log.application) == {corrections.APPLIED}
    sep = lambda r: r.residuals.query("rule == 'R3' and building == '*'").diff_m2.iloc[0]  # noqa: E731
    assert sep(before) > 15000 and sep(after) < 300          # restent les petits écarts d'étage (S4, S5)


@needs_tgm
def test_review_v2_reorients_subjects(loree):
    # Revue L'Orée (docs/revue-ia, version 2) : S1 et S2 gardent les propositions, S4–S6 sans correction.
    pytest.importorskip("shapely")
    from sia_check import config, corrections, geom, sujets, viewer
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    df, res = loree
    doc = viewer.build(res, cfg, geom.read_cache(LOREE), "L OREE", [str(TGM)])
    rv = sujets.load_review(Path(__file__).resolve().parents[1] / "docs/revue-ia/L_OREE_2026-10-22.verdicts.json")
    out = sujets.apply_review(sujets.build(doc, cfg), rv, {s["id"] for s in doc["surfaces"]})
    assert out["revue"]["problemes"] == [] and out["revue"]["decisions"] == 31 and not out["revue"]["points_inconnus"]
    by = {s["id"]: s for s in out["subjects"]}
    assert len(by["S1"]["corrections"]) == 9 and len(by["S2"]["corrections"]) == 22
    assert len(by["S6"]["points"]) == 71 and not by["S6"]["corrections"]
    j = sujets.journal(out)
    for c in j["corrections"]:
        c["statut"] = "acceptée"
    applied, before, after = corrections.run(df, j, cfg)
    assert set(applied.log.application) == {corrections.APPLIED}
    assert corrections.counts(after)["ERREUR"] < corrections.counts(before)["ERREUR"]


@needs_tgm
def test_mdp_proposals_on_loree(loree):
    from sia_check import config, propositions
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    _, res = loree
    by = {s["id"]: s for s in propositions.build(res, cfg)}
    aff = {c["id"]: c["affectation"] for c in by["affectation"]["corrections"]}
    assert aff["affectation-Gaine-lot"] == "Gaines" and aff["affectation-Chambre-lot"] == "Logement general"
    assert [c["id"] for c in by["typologies"]["corrections"]] == [f"typologie-{n}ch" for n in range(5)]   # sans le lot à −1
    assert len(by["formes"]["corrections"]) == 9


@needs_tgm
def test_mdp_ready_for_power_bi(loree, tmp_path):
    # Points relevés en comparant au MDP 1778 fait à la main : étages au code du gabarit, pas de SP à 0 m²,
    # chambres −1 laissées vides, ID projet sur toutes les lignes.
    import warnings

    import openpyxl

    from sia_check import config, mdp
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    _, res = loree
    out = mdp.write(res, cfg, tmp_path / "MDP.xlsx", project_id=1778, project_name="L'Orée")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(out, data_only=True)
    brutes = list(wb["SurfacesBrutesVolumeSIA416"].iter_rows(min_row=2, values_only=True))
    assert {r[0] for r in brutes} == {1778} and all(r[10] for r in brutes if r[1] == "SP")
    assert {(r[3], r[4]) for r in brutes if r[3] == "TOI"} == {("TOI", "Hors-sol")}
    assert "Toiture" not in {r[3] for ws in ("SurfacesEnveloppe", "SurfacesBrutesVolumeSIA416")
                             for r in wb[ws].iter_rows(min_row=2, values_only=True)}
    lots = {r[3]: r[10] for r in wb["ListeLogements"].iter_rows(min_row=2, values_only=True)}
    assert lots["I18_N01_Apart-03"] is None and min(v for v in lots.values() if v is not None) == 0
    info = {r[1]: r[2] for r in wb["InfosProjet"].iter_rows(min_row=2, values_only=True)}
    assert info["Nbr de chambres"] == 620 and info["Nom Projet HRS"] == "L'Orée"


@needs_tgm
def test_sec_implicit_after_corrections(loree):
    # Treegram ne crée pas la SEC : une fois toitures et abords corrigés, le résidu R3 des balcons est implicite.
    from sia_check import config, corrections
    cfg = config.load(Path(__file__).resolve().parents[1] / "config.yaml")
    df, _ = loree
    journal = json.loads((Path(__file__).resolve().parents[1] / "docs/corrections/L_OREE_exemple.json").read_text(encoding="utf-8"))
    _, _, after = corrections.run(df, journal, cfg)
    r3 = after.residuals[(after.residuals.rule == "R3") & (after.residuals.level != "*")]
    rel = r3.diff_m2 / r3.expected
    assert (r3.verdict == "INFO").sum() >= 10 and (rel[r3.verdict == "INFO"] <= 0.35).all()   # SEC implicite
    assert (rel[r3.verdict == "ERREUR"] > 0.35).all()          # abords, attiques, I17 S01 : vrais sujets
