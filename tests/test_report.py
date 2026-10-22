"""Rapports xlsx / html / json sur le jeu synthétique."""

import json

import openpyxl
import pytest

from sia_check import report, rules


@pytest.fixture
def written(result, cfg, tmp_path):
    rep = report.build(result, cfg, "Projet test/1", ["synth.csv"], date="2026-10-21")
    return rep, report.write_all(rep, tmp_path / "out")


def test_file_names(written):
    _, paths = written
    assert [p.name for p in paths] == [f"Projet_test_1_2026-10-21.{e}" for e in ("xlsx", "html", "json")]
    assert all(p.is_file() and p.stat().st_size > 0 for p in paths)


def test_xlsx_sheets(written):
    rep, paths = written
    wb = openpyxl.load_workbook(paths[0])
    assert wb.sheetnames == ["Résidus", "Surfaces", "Codes inconnus", "Synthèse", "Contrôles"]
    header = [c.value for c in wb["Résidus"][1]]
    for col in ("Bâtiment", "Niveau", "Règle", "Attendu (m²)", "Obtenu (m²)", "Écart (m²)", "Écart (%)", "Verdict"):
        assert col in header
    assert wb["Résidus"].max_row - 1 == len(rep.residuals)
    assert wb["Surfaces"].max_row - 1 == len(rep.result.surfaces) + len(rep.result.unknown)
    assert wb["Codes inconnus"].max_row - 1 == 2


def test_residuals_sorted_by_gap(written):
    rep, _ = written
    gaps = rep.residuals.diff_m2.abs().tolist()
    assert gaps == sorted(gaps, reverse=True)
    assert "OK" not in set(rep.residuals.verdict)


def test_synthesis_counts(written):
    rep, _ = written
    cat = rep.categories.set_index("Catégorie")
    assert cat.loc["SP", "Totaux de niveau"] == 8      # OK ×2, ERR ×5, TEC ×1
    assert cat.loc["SUT", "Statut"] == "contrôlée"
    assert cat.loc["XYZ", "Statut"] == "inconnue"
    assert rep.buildings.loc["Total", "SP"] == pytest.approx(740)


def test_json_for_n8n(written):
    rep, paths = written
    doc = json.loads(paths[2].read_text(encoding="utf-8"))
    assert doc["exit_code"] == 1 and doc["project"] == "Projet test/1"
    assert len(doc["residuals"]) == len(rep.residuals)
    assert {u["category"] for u in doc["unknown_codes"]} == {"XYZ", "(vide)"}
    assert doc["counts"]["ERREUR"] > 0


def test_html(written):
    _, paths = written
    html = paths[1].read_text(encoding="utf-8")
    assert "Projet test/1" in html and "#B1005D" in html and "repeating-linear-gradient" in html
    assert "Bâtiment ERR" in html


def test_clean_model_report_is_empty(synth, cfg, tmp_path):
    res = rules.run(synth[synth.building == "OK"], cfg)
    rep = report.build(res, cfg, "propre", ["x"], date="2026-10-21")
    paths = report.write_all(rep, tmp_path)
    assert json.loads(paths[2].read_text(encoding="utf-8"))["residuals"] == []
    assert openpyxl.load_workbook(paths[0])["Résidus"].max_row == 1
    assert "Aucun résidu." in paths[1].read_text(encoding="utf-8")
