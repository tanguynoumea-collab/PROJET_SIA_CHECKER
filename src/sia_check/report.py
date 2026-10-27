"""Rapports xlsx, html et json à partir du résultat des règles (rules.Result).

Fichiers produits : <out>/<projet>_<date>.xlsx | .html | .json
- xlsx : onglets Résidus, Surfaces, Codes inconnus, Synthèse, Contrôles (tous les contrôles, OK compris)
- html : même contenu, lisible sans Excel ; résidus triés par écart décroissant ; tableau par niveau
         aux couleurs de la légende SIA 416
- json : résidus en machine-readable pour n8n
"""

from __future__ import annotations

import datetime as dt
import json
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from jinja2 import Environment, PackageLoader, select_autoescape

from . import __version__, rules
from .config import Config

# Légende officielle SIA 416 (figures de la norme). Surfaces externes (§3) : couleur du pendant intérieur, hachurée.
# SUT (catégorie Treegram comptée dans la SU, règle P1) : orangé, entre SUP et SUS.
SIA_COLORS = {"SP": "#B1005D", "SUP": "#D5003B", "SUT": "#E8601C", "SUS": "#F0A900", "SD": "#F2E900", "SI": "#0C9DDA",
              "SEP": "#B1005D", "SEU": "#D5003B", "SED": "#F2E900", "SEI": "#0C9DDA",
              "SA": "#5BA845",   # abords (§1.2) : surfaces créées par une découpe vers SA
              "SFN": "#A8A29E", "SRV": "#9ED8E6"}   # enveloppe : hors légende SIA 416 (DESIGN_PLAN, Évolution 3.6)
HATCHED = {"SEP", "SEU", "SED", "SEI", "SEN", "SEC"}
LEVEL_COLUMNS = ["SP", "SUP", "SUT", "SUS", "SD", "SI", "SEP", "SEU", "SED"]

RESIDUAL_LABELS = {
    "building": "Bâtiment", "level": "Niveau", "lot": "Lot", "rule": "Règle", "category": "Catégorie",
    "expected": "Attendu (m²)", "obtained": "Obtenu (m²)", "diff_m2": "Écart (m²)", "diff_pct": "Écart (%)",
    "verdict": "Verdict", "message": "Constat", "hint": "À vérifier",
}
SURFACE_LABELS = {
    "id": "Identifiant (GlobalId / élément)", "building": "Bâtiment", "level": "Niveau", "category": "Catégorie",
    "code": "Code normalisé", "status": "Statut", "kind": "Type", "area": "Aire (m²)",
    "room_type": "Sous-type", "lot": "Lot", "source": "Source",
}
RATIO_LABELS = {
    "building": "Portée", "id": "Ratio", "family": "Famille", "label": "Libellé", "value": "Valeur", "unit": "Unité",
    "ref_min": "Repère min", "ref_max": "Repère max", "status": "Statut", "num_m2": "Numérateur", "den_m2": "Dénominateur",
    "note": "Note",
}
KIND_LABELS = {"leaf": "élémentaire", "node": "total de niveau"}


@dataclass
class Report:
    project: str
    date: str
    sources: list[str]
    cfg: Config
    result: rules.Result
    residuals: pd.DataFrame = field(init=False)     # contrôles hors OK, écart décroissant
    surfaces: pd.DataFrame = field(init=False)      # toutes les surfaces lues, statut compris
    categories: pd.DataFrame = field(init=False)    # synthèse par catégorie
    buildings: pd.DataFrame = field(init=False)     # synthèse bâtiment × catégorie
    levels: pd.DataFrame = field(init=False)        # tableau par niveau

    def __post_init__(self):
        res = self.result.residuals
        # Points d'attention : contrôles hors OK, sans les sommes que les niveaux détaillent déjà.
        nok = res[(res.verdict != "OK") & ~rules.covered_aggregates(res)]
        self.residuals = nok.reindex(nok.diff_m2.abs().sort_values(ascending=False).index)
        self.surfaces = _surfaces(self.result)
        self.categories = _category_summary(self.result, self.cfg)
        self.buildings = _building_summary(self.result, self.cfg)
        self.levels = _level_table(self.result, self.cfg)

    @property
    def stem(self) -> str:
        return f"{_safe(self.project)}_{self.date}"

    @property
    def counts(self) -> dict[str, int]:
        vc = self.residuals.verdict.value_counts()
        vc["OK"] = int((self.result.residuals.verdict == "OK").sum())
        return {v: int(vc.get(v, 0)) for v in rules.VERDICTS}


def build(result: rules.Result, cfg: Config, project: str, sources: list[str], date: str | None = None) -> Report:
    return Report(project, date or dt.date.today().isoformat(), [str(s) for s in sources], cfg, result)


def write_all(rep: Report, out_dir: str | Path) -> list[Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    paths = [out / f"{rep.stem}.{ext}" for ext in ("xlsx", "html", "json")]
    write_xlsx(rep, paths[0])
    write_html(rep, paths[1])
    write_json(rep, paths[2])
    return paths


# --------------------------------------------------------------------------- tables

def _safe(name: str) -> str:
    return re.sub(r"[^\w\-]+", "_", name, flags=re.UNICODE).strip("_") or "projet"


def _category_order(cfg: Config, codes) -> list[str]:
    known = [c for c in cfg.categories if c in set(codes)]
    return known + sorted(set(codes) - set(known))


def _surfaces(result: rules.Result) -> pd.DataFrame:
    s = result.surfaces.assign(status=result.surfaces.ignored.map({True: "ignorée", False: "contrôlée"}))
    u = result.unknown.assign(code=None, status="inconnue")
    df = pd.concat([s, u], ignore_index=True)
    df["kind"] = df["kind"].map(KIND_LABELS).fillna(df["kind"])
    return df.reindex(columns=list(SURFACE_LABELS)).sort_values(["building", "level", "category", "id"], na_position="last")


def _category_summary(result: rules.Result, cfg: Config) -> pd.DataFrame:
    """Une ligne par catégorie : de quoi rapprocher les compteurs du viewer Treegram (ex. 80 SP)."""
    s = result.surfaces
    tot = rules.level_totals(s, include_ignored=True)
    rows = []
    for code in _category_order(cfg, list(tot.columns)):
        cs = s[s.code == code]
        rows.append({
            "Catégorie": code,
            "Famille": cfg.categories.get(code, {}).get("family", ""),
            "Statut": "ignorée" if cs.ignored.all() else "contrôlée",
            "Niveaux": int(tot[code].notna().sum()),
            "Totaux de niveau": int((cs.kind == "node").sum()),
            "Surfaces élémentaires": int((cs.kind == "leaf").sum()),
            "Aire (m²)": float(tot[code].sum()),
        })
    for code, g in result.unknown.groupby(result.unknown.category.replace("", "(vide)")):
        rows.append({"Catégorie": code, "Famille": "", "Statut": "inconnue",
                     "Niveaux": int(g.groupby(["building", "level"]).ngroups),
                     "Totaux de niveau": int((g.kind == "node").sum()),
                     "Surfaces élémentaires": int((g.kind == "leaf").sum()),
                     "Aire (m²)": float(g.area.sum())})
    return pd.DataFrame(rows)


def _building_summary(result: rules.Result, cfg: Config) -> pd.DataFrame:
    tot = rules.level_totals(result.surfaces, include_ignored=True)
    if tot.empty:
        return pd.DataFrame()
    b = tot.groupby(level="building").sum(min_count=1)
    b.loc["Total"] = b.sum(min_count=1)
    b = b[_category_order(cfg, list(b.columns))]
    b.index.name = "Bâtiment"
    return b


def _level_table(result: rules.Result, cfg: Config) -> pd.DataFrame:
    tot = result.totals
    if tot.empty:
        return pd.DataFrame()
    t = tot.reindex(columns=LEVEL_COLUMNS)
    res = result.residuals
    lv = res[(res.level != rules.ALL)]
    r1 = lv[lv.category == "SP"].set_index(["building", "level"])
    r3 = lv[lv.category == "SEP"].set_index(["building", "level"])
    fails = lv[lv.verdict != "OK"].groupby(["building", "level"])
    t["R1 écart"] = r1["diff_m2"].groupby(level=[0, 1]).first().reindex(t.index)
    t["R1 verdict"] = r1["verdict"].groupby(level=[0, 1]).first().reindex(t.index)
    t["R3 écart"] = r3[r3.rule.isin(["R3", "R4"])]["diff_m2"].groupby(level=[0, 1]).first().reindex(t.index)
    t["R3 verdict"] = r3[r3.rule.isin(["R3", "R4"])]["verdict"].groupby(level=[0, 1]).first().reindex(t.index)
    t["Alertes"] = fails["rule"].agg(lambda r: ", ".join(sorted(set(r) - {"R1", "R3"}))).reindex(t.index).fillna("")
    t["_rank"] = [rules.level_rank(l, cfg) for _, l in t.index]
    t = t.reset_index().sort_values(["building", "_rank", "level"], na_position="last").drop(columns="_rank")
    return t


# --------------------------------------------------------------------------- xlsx

def write_xlsx(rep: Report, path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    verdict_fill = {"ERREUR": "F8D7DA", "ALERTE": "FFE8B3", "INFO": "D6E9F8", "OK": "D8F0DE"}
    sheets = {
        "Résidus": rep.residuals.rename(columns=RESIDUAL_LABELS),
        "Surfaces": rep.surfaces.rename(columns=SURFACE_LABELS),
        "Codes inconnus": rep.surfaces[rep.surfaces.status == "inconnue"]
        .drop(columns=["code", "status"]).rename(columns=SURFACE_LABELS),
        "Synthèse": rep.categories,
        "Contrôles": rep.result.residuals.rename(columns=RESIDUAL_LABELS),
        "Ratios": _ratios_frame(rep).rename(columns=RATIO_LABELS),
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        start = len(rep.categories) + 3
        if not rep.buildings.empty:
            ws = xw.sheets["Synthèse"]
            ws.cell(row=start, column=1, value="Aire par bâtiment et par catégorie (m²)").font = Font(bold=True)
            rep.buildings.to_excel(xw, sheet_name="Synthèse", startrow=start)

        for name, ws in xw.sheets.items():
            ws.freeze_panes = "A2"
            headers = {}
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="EDEDED")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                headers[cell.column] = str(cell.value or "")
            for col in ws.iter_cols(min_row=2):
                if not col:  # onglet sans ligne de données (ex. aucun code inconnu)
                    break
                h = headers.get(col[0].column, "")
                for c in col:
                    if isinstance(c.value, float):
                        c.number_format = "0.0" if "%" in h else "#,##0.00"
                    if h == "Verdict" and c.value in verdict_fill:
                        c.fill = PatternFill("solid", fgColor=verdict_fill[c.value])
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 9), 60)


# --------------------------------------------------------------------------- html

def _fmt(v, digits=2) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if isinstance(v, (int, float)):
        return f"{v:,.{digits}f}".replace(",", " ")
    return str(v)


def _tint(hex_color: str, alpha: float) -> str:
    h = hex_color.lstrip("#")
    r, g, b = (int(h[i : i + 2], 16) for i in (0, 2, 4))
    return f"rgba({r},{g},{b},{alpha})"


def write_html(rep: Report, path: Path) -> None:
    env = Environment(loader=PackageLoader("sia_check", "templates"), autoescape=select_autoescape(["html"]))
    env.filters["num"] = _fmt
    swatches = {c: {"color": SIA_COLORS.get(c, "#888"), "tint": _tint(SIA_COLORS.get(c, "#888888"), 0.10),
                    "hatched": c in HATCHED} for c in LEVEL_COLUMNS}
    by_building = {b: g.to_dict("records") for b, g in rep.levels.groupby("building", sort=True)} \
        if not rep.levels.empty else {}
    html = env.get_template("report.html.j2").render(
        rep=rep, version=__version__, counts=rep.counts, exit_code=rep.result.exit_code,
        residuals=rep.residuals.to_dict("records"), levels=by_building, level_columns=LEVEL_COLUMNS,
        swatches=swatches, categories=rep.categories.to_dict("records"),
        buildings=rep.buildings.reset_index().to_dict("records") if not rep.buildings.empty else [],
        building_columns=list(rep.buildings.columns),
        unknown=rep.surfaces[rep.surfaces.status == "inconnue"].to_dict("records"),
        surfaces=rep.surfaces.to_dict("records"), kind_labels=KIND_LABELS,
        ratios=_ratios_frame(rep).to_dict("records"),
    )
    path.write_text(html, encoding="utf-8")


# --------------------------------------------------------------------------- json

def _records(df: pd.DataFrame) -> list[dict]:
    return json.loads(df.round(3).to_json(orient="records", force_ascii=False))


def _ratios_frame(rep: Report) -> pd.DataFrame:
    """Ratios informatifs (ratios.py) ; portée « * » = projet entier."""
    r = rep.result.ratios
    return r if r is not None else pd.DataFrame(columns=list(RATIO_LABELS))


def write_json(rep: Report, path: Path) -> None:
    unknown = rep.categories[rep.categories["Statut"] == "inconnue"]
    doc = {
        "tool": "sia-check",
        "version": __version__,
        "project": rep.project,
        "date": rep.date,
        "sources": rep.sources,
        "exit_code": rep.result.exit_code,
        "tolerance": {"abs_m2": rep.cfg.abs_m2, "rel": rep.cfg.rel},
        "custom_rules": [{"id": r.id, "title": r.title, "description": r.description} for r in rep.cfg.custom_rules],
        "counts": rep.counts,
        "residuals": _records(rep.residuals),
        "ratios": _records(_ratios_frame(rep)),
        "unknown_codes": [{"category": r["Catégorie"], "count": r["Surfaces élémentaires"] + r["Totaux de niveau"],
                           "area_m2": round(r["Aire (m²)"], 3)} for r in unknown.to_dict("records")],
    }
    path.write_text(json.dumps(doc, ensure_ascii=False, indent=2), encoding="utf-8")
