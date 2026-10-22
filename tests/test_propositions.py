"""Propositions MDP (propositions.py), « renseigner » / « Inconnu » (corrections.py), Excel de correction et SA (treegram.py)."""

import warnings

import openpyxl
import pandas as pd
import pytest

from sia_check import corrections, mdp, propositions, rules, treegram


def surfaces() -> pd.DataFrame:
    rows = [
        # id, catégorie, aire, sorte, type, lot, chambres, dx, dy
        ("A_N01_SP#total", "SP", 200.0, "node", None, None, None, 10.0, 20.0),
        ("A_N01_SEP#total", "SEP", 60.0, "node", None, None, None, None, None),
        ("ch-1", "SUP", 20.0, "leaf", "Chambre", "A_N01_Apart-1", None, None, None),
        ("gaine-1", "SI", 1.0, "leaf", "Gaine", "A_N01_Apart-1", None, None, None),
        ("esc-1", "SD", 15.0, "leaf", "Escalier", None, None, None, None),
        ("snh-1", "SNH", 60.0, "leaf", None, "A_N01_Apart-1", 2, None, None),
        ("snh-2", "SNH", 40.0, "leaf", None, "A_N01_Apart-2", 1, None, None),
    ]
    return pd.DataFrame([{"id": i, "building": "A", "level": "N01", "category": c, "area": a, "kind": k, "room_type": t,
                          "lot": l, "bedrooms": b, "dx": dx, "dy": dy} for i, c, a, k, t, l, b, dx, dy in rows])


@pytest.fixture
def subjects(cfg):
    return {s["id"]: s for s in propositions.build(rules.run(surfaces(), cfg), cfg)}


def test_affectation_by_room_type_and_housing(subjects):
    by = {c["id"]: c for c in subjects["affectation"]["corrections"]}
    assert by["affectation-Chambre-lot"]["affectation"] == "Logement general"
    assert by["affectation-Gaine-lot"]["affectation"] == "Gaines"          # une gaine reste une gaine, même en logement
    assert by["affectation-Escalier"]["cible"]["selection"]["avec_lot"] is False
    assert all(c["statut"] == "proposée" and c["origine"] == "règle" for c in by.values())


def test_typologies_and_shape(subjects):
    typo = {c["id"]: c["champs"] for c in subjects["typologies"]["corrections"]}
    assert typo["typologie-2ch"] == {"Nom typologie": "3.5P", "Nbr de pieces (VD)": 3.5, "Nbr de pieces (GE)": 4}
    assert subjects["formes"]["corrections"][0]["champs"] == {"Forme du bâtiment": "Barre"}   # 10 × 20, pleine
    assert {s["action"] for s in subjects.values()} == {"proposer"}


def test_accepted_and_unknown_reach_the_mdp(cfg, subjects, tmp_path):
    cs = [c for s in subjects.values() for c in s["corrections"]]
    for c in cs:
        c["statut"] = "rejetée" if c["id"] == "typologie-1ch" else "acceptée"
    applied, _, after = corrections.run(surfaces(), {"corrections": cs}, cfg)
    assert set(applied.log.application) == {corrections.APPLIED}
    out = mdp.write(after, cfg, tmp_path / "m.xlsx", infos=applied.infos)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(out)
    lots = {r[3]: (r[5], r[8], r[9]) for r in wb["ListeLogements"].iter_rows(min_row=2, values_only=True)}
    assert lots["A_N01_Apart-1"] == ("3.5P", 4, 3.5)
    assert lots["A_N01_Apart-2"] == ("Inconnu", "Inconnu", "Inconnu")    # proposition refusée
    nettes = {r[3] for r in wb["SurfacesNettesSIA416"].iter_rows(min_row=2, values_only=True)}
    assert nettes >= {"Logement general", "Gaines", "Cage d'escaliers"}
    assert [r[6] for r in wb["InfosBatiment"].iter_rows(min_row=2, values_only=True)] == ["Barre"]


def test_rejected_affectation_writes_unknown(cfg):
    j = {"corrections": [{"id": "a", "op": "affecter", "statut": "rejetée", "affectation": "Gaines",
                          "cible": {"selection": {"room_type": "Gaine"}}}]}
    a = corrections.apply(surfaces(), j)
    assert a.surfaces.set_index("id").loc["gaine-1", "affectation"] == "Inconnu"
    assert a.log.iloc[0].raison.startswith("proposition refusée")


def test_correction_workbook_and_abords_rows(cfg, tmp_path):
    j = {"corrections": [
        {"id": "d", "op": "découper", "statut": "acceptée", "cible": {"ids": ["A_N01_SEP#total"]}, "vers": "SA", "aire": 25.0,
         "explication": "Dalle sur sous-sol."},
        {"id": "x", "op": "exclure", "statut": "acceptée", "cible": {"ids": ["esc-1"]}, "explication": "Doublon."}]}
    applied, _, after = corrections.run(surfaces(), j, cfg)
    assert ("SA", 25.0) in {(r["Catégorie SIA"], r["Area"]) for r in treegram.brutes(after, cfg).to_dict("records")}
    nets, gross, log = treegram.corrections_frames(applied, j, cfg)
    sep = gross[gross["Object.Name"] == "SEP"].iloc[0]
    assert (sep["Area"], sep["Aire corrigée"], sep["Correction"]) == (60.0, 35.0, "d (découper → SA)")
    sa = gross[gross["Catégorie corrigée"] == "SA"].iloc[0]
    assert pd.isna(sa["Area"]) and sa["Aire corrigée"] == 25.0 and sa["Explication"] == "Dalle sur sous-sol."
    esc = nets[nets["Object.Name"] == "esc-1"].iloc[0]
    assert (esc["Exclue"], esc["Aire corrigée"], esc["Catégorie SIA"]) == ("oui", 0.0, "SD")
    out = treegram.write_corrections(applied, j, cfg, tmp_path / "c.xlsx")
    assert pd.ExcelFile(out).sheet_names == ["SurfacesNettes", "SurfacesBrutes", "Journal"]
    assert list(log.explication) == ["Dalle sur sous-sol.", "Doublon."]
